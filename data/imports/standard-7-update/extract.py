import json, re, pathlib
from docx import Document
source = pathlib.Path(r'C:\Users\ishma\Downloads\Questionnaires for Lae\Environment\Standard_7_Physical_Infrastructure_Environment_Management_PNG.docx')
d = Document(source)
criteria=[]; compliances=[]; evidence=[]; cn=None; co=None; ignored=[]
for ti in (1,2):
 for row in d.tables[ti].rows:
  cells=list(dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip()))
  if not cells: continue
  first=cells[0]
  m=re.fullmatch(r'7\.(\d+)\.(\d+)\.?',first)
  if m:
   num=first.rstrip('.')
   if cn=='7.14' and num=='7.4.1': num='7.14.1'
   assert num.rsplit('.',1)[0]==cn,(cn,num)
   co=num; compliances.append([cn,num,cells[1]]); continue
  m=re.match(r'^7\.\s*(\d+)\.?\s*(.*)$',first)
  if m and (len(cells)>1 or m.group(2)) and not first.startswith('7.  Planned'):
   cn='7.'+m.group(1); co=None; criteria.append([cn,m.group(2) or cells[1]]); continue
  if co:
   m=re.match(r'^(\d+)\.\s*(.*)$',first,re.S)
   if m:
    body=m.group(2).strip() or (cells[1] if len(cells)>1 else '')
    assert body,(co,cells)
    evidence.append([co,m.group(1),body]); continue
  ignored.append(cells)
assert len(criteria)==14,criteria
assert len({c[1] for c in compliances})==len(compliances)
assert len({tuple(e[:2]) for e in evidence})==len(evidence)
assert all(any(e[0]==c[1] for e in evidence) for c in compliances)
title='\n'.join(dict.fromkeys(c.text.strip() for r in d.tables[0].rows for c in r.cells if c.text.strip()))
data={'header':title,'criteria':criteria,'compliances':compliances,'evidence':evidence,'ignored':ignored}
pathlib.Path(__file__).with_name('extracted.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'criteria':len(criteria),'compliances':len(compliances),'evidence':len(evidence),'ignored':ignored},ensure_ascii=True))
